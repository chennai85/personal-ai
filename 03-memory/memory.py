# /// script
# requires-python = ">=3.11"
# dependencies = ["numpy>=1.26", "pypdf>=4.0", "python-docx>=1.1"]
# ///
"""
memory.py: my personal memory, on the Mac Studio.

    uv run memory.py ingest            read the vault, index what's new or changed
    uv run memory.py search "QUERY"    find passages (meaning + exact words)
    uv run memory.py ask "QUESTION"    the main model answers from my files, with sources
    uv run memory.py status            what's indexed, and is everything working
    uv run memory.py rebuild           delete the index and build it again

The vault (~/vault) holds my files. The index (~/vault/.index/memory.db)
is built from them and can be deleted at any time: the files are the memory.

Settings come from 01-environment/standard.sh (the single source of truth).
"""

import argparse
import email
import email.policy
import hashlib
import html
import json
import os
import re
import sqlite3
import sys
import time
import urllib.request
from pathlib import Path

import numpy as np

HERE = Path(__file__).resolve().parent
STANDARD = HERE.parent / "01-environment" / "standard.sh"
SUPPORTED = {".md", ".txt", ".pdf", ".docx", ".eml", ".html", ".htm"}
# Bump when the way files are read or split changes: every file is then re-read.
CHUNKER_VERSION = "2"   # v2: skip blank template lines, strip repeated PDF headers/footers


# ---------------------------------------------------------------------
# Settings: read from standard.sh, environment can override
# ---------------------------------------------------------------------
def load_settings():
    vals = {}
    if STANDARD.exists():
        for line in STANDARD.read_text().splitlines():
            m = re.match(r'^([A-Z_]+)=("([^"]*)"|(\S+))', line.strip())
            if m:
                vals[m.group(1)] = m.group(3) if m.group(3) is not None else m.group(4)
    home = str(Path.home())

    def get(key, default):
        v = os.environ.get(key) or vals.get(key) or default
        v = re.sub(r"\$\{(\w+):-([^}]*)\}", lambda m: os.environ.get(m.group(1)) or m.group(2), v)
        v = v.replace("$HOME", home)
        for k, kv in vals.items():          # expand $VAULT_DIR etc.
            v = v.replace(f"${k}", os.environ.get(k) or kv.replace("$HOME", home))
        return v

    s = {
        "vault": Path(get("VAULT_DIR", f"{home}/vault")),
        "embed_model": get("MODEL_EMBED", "qwen3-embedding:0.6b"),
        "main_model": get("MODEL_MAIN", "qwen3.6:35b-a3b"),
        "context": int(get("OLLAMA_CONTEXT_LENGTH", "32768")),
        "model_url": get("MODEL_BASE_URL", "http://localhost:11434").rstrip("/"),
        "chunk_words": int(get("MEMORY_CHUNK_WORDS", "500")),
        "overlap": int(get("MEMORY_CHUNK_OVERLAP", "80")),
    }
    idx = os.environ.get("VAULT_INDEX")
    s["index"] = Path(idx) if idx else s["vault"] / ".index" / "memory.db"
    return s


S = load_settings()


# ---------------------------------------------------------------------
# Embeddings (Ollama). Each model family wants its own prefixes.
# ---------------------------------------------------------------------
def prefixes(model):
    m = model.lower()
    if m.startswith("qwen3-embedding"):
        return ("Instruct: Given a question, retrieve passages from my personal files "
                "that answer it\nQuery: ", "")
    if m.startswith("embeddinggemma"):
        return ("task: search result | query: ", "title: none | text: ")
    if m.startswith("nomic-embed"):
        return ("search_query: ", "search_document: ")
    return ("", "")


def embed(texts, kind="doc"):
    """Return L2-normalised float32 vectors, one per text."""
    q_pre, d_pre = prefixes(S["embed_model"])
    pre = q_pre if kind == "query" else d_pre
    out = []
    for i in range(0, len(texts), 32):
        batch = [pre + t for t in texts[i:i + 32]]
        req = urllib.request.Request(
            f"{S['model_url']}/api/embed",
            data=json.dumps({"model": S["embed_model"], "input": batch}).encode(),
            headers={"Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(req, timeout=300) as r:
                out.extend(json.loads(r.read())["embeddings"])
        except Exception as e:
            sys.exit(f"STOP: could not reach the embedding model {S['embed_model']} at "
                     f"{S['model_url']} ({e.__class__.__name__}: {e}).\n"
                     f"Is the Studio awake? Was it pulled? (bash 02-local-model/server_setup.sh)")
    v = np.asarray(out, dtype=np.float32)
    v /= np.linalg.norm(v, axis=1, keepdims=True) + 1e-12
    return v


# ---------------------------------------------------------------------
# The index
# ---------------------------------------------------------------------
def open_db():
    S["index"].parent.mkdir(parents=True, exist_ok=True)
    os.chmod(S["index"].parent, 0o700)
    con = sqlite3.connect(S["index"])
    con.executescript("""
        PRAGMA journal_mode=WAL;
        CREATE TABLE IF NOT EXISTS meta(key TEXT PRIMARY KEY, value TEXT);
        CREATE TABLE IF NOT EXISTS files(
            path TEXT PRIMARY KEY, sha256 TEXT, size INTEGER, kind TEXT,
            status TEXT, note TEXT, chunks INTEGER, ingested_at TEXT);
        CREATE TABLE IF NOT EXISTS chunks(
            id INTEGER PRIMARY KEY, path TEXT, ord INTEGER, page INTEGER,
            text TEXT, embedding BLOB);
        CREATE INDEX IF NOT EXISTS chunks_path ON chunks(path);
    """)
    try:
        con.execute("CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts "
                    "USING fts5(text, tokenize='porter unicode61')")
    except sqlite3.OperationalError as e:
        sys.exit(f"STOP: this Python's SQLite has no FTS5 (keyword search): {e}")
    return con


def meta(con, key, value=None):
    if value is None:
        r = con.execute("SELECT value FROM meta WHERE key=?", (key,)).fetchone()
        return r[0] if r else None
    con.execute("INSERT OR REPLACE INTO meta VALUES (?,?)", (key, str(value)))


def drop_file(con, path):
    ids = [r[0] for r in con.execute("SELECT id FROM chunks WHERE path=?", (path,))]
    con.executemany("DELETE FROM chunks_fts WHERE rowid=?", [(i,) for i in ids])
    con.execute("DELETE FROM chunks WHERE path=?", (path,))
    con.execute("DELETE FROM files WHERE path=?", (path,))


# ---------------------------------------------------------------------
# Reading files: each returns [(page, text), ...]
# ---------------------------------------------------------------------
def read_pdf(p):
    from pypdf import PdfReader
    pages = [(i + 1, page.extract_text() or "") for i, page in enumerate(PdfReader(p).pages)]
    return strip_repeated_lines(pages)


def strip_repeated_lines(pages, share=0.6):
    """Drop header/footer lines that repeat on most pages ("Page 3", the same title).

    Repeated lines make every page look alike to search, so the page that
    actually answers the question gets buried. Digits are ignored when comparing,
    so "Page 3 of 9" and "Page 4 of 9" count as the same line.
    """
    if len(pages) < 3:
        return pages
    norm = lambda line: re.sub(r"\d+", "#", line.strip().lower())
    counts = {}
    for _, text in pages:
        for key in {norm(l) for l in text.splitlines() if l.strip()}:
            counts[key] = counts.get(key, 0) + 1
    repeated = {k for k, n in counts.items() if n >= share * len(pages)}
    return [(pg, "\n".join(l for l in text.splitlines() if norm(l) not in repeated))
            for pg, text in pages]


def read_docx(p):
    import docx
    d = docx.Document(p)
    parts = [para.text for para in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return [(1, "\n".join(parts))]


def read_eml(p):
    msg = email.message_from_bytes(Path(p).read_bytes(), policy=email.policy.default)
    head = (f"From: {msg['from']}\nTo: {msg['to']}\nDate: {msg['date']}\n"
            f"Subject: {msg['subject']}\n\n")
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    if body and body.get_content_type() == "text/html":
        text = strip_html(text)
    attach = [a.get_filename() for a in msg.iter_attachments() if a.get_filename()]
    if attach:
        text += "\n\nAttachments: " + ", ".join(attach)
    return [(1, head + text)]


def strip_html(s):
    s = re.sub(r"(?is)<(script|style).*?</\1>", " ", s)
    s = re.sub(r"(?s)<br\s*/?>|</p>|</div>|</li>|</h\d>", "\n", s)
    return html.unescape(re.sub(r"(?s)<[^>]+>", " ", s))


def read_file(p):
    ext = p.suffix.lower()
    if ext == ".pdf":
        return read_pdf(p)
    if ext == ".docx":
        return read_docx(p)
    if ext == ".eml":
        return read_eml(p)
    if ext in (".html", ".htm"):
        return [(1, strip_html(p.read_text(errors="replace")))]
    return [(1, skip_blanks(p.read_text(errors="replace")))]


def skip_blanks(text):
    """Leave out the parts of a template I haven't filled in yet.

    Blank fields contain ⟪ ⟫. In a file that still has blanks, also leave out the
    front matter (its "about:" line would match searches), table headers with no
    rows, and headings with nothing under them. A filled-in file is indexed as written.
    """
    if "⟪" not in text:
        return text
    lines = text.splitlines()
    if lines and lines[0].strip() == "---":
        end = next((i for i, l in enumerate(lines[1:], 1) if l.strip() == "---"), 0)
        lines = lines[end + 1:]
    is_sep = lambda l: re.fullmatch(r"\s*\|[-|\s:]*\|\s*", l) is not None
    lines = [l for l in lines if "⟪" not in l and not is_sep(l) and l.strip()]
    is_row = lambda l: l.lstrip().startswith("|")
    is_head = lambda l: l.lstrip().startswith("#")
    # a table header with no rows under it
    lines = [l for i, l in enumerate(lines)
             if not (is_row(l) and (i + 1 == len(lines) or not is_row(lines[i + 1]))
                     and (i == 0 or not is_row(lines[i - 1])))]
    # a heading with nothing under it (until the next heading of the same or higher level)
    level = lambda l: len(l.lstrip()) - len(l.lstrip().lstrip("#"))
    keep = []
    for i, l in enumerate(lines):
        if is_head(l):
            nxt = lines[i + 1] if i + 1 < len(lines) else None
            if nxt is None or (is_head(nxt) and level(nxt) <= level(l)):
                continue
        keep.append(l)
    return "\n".join(keep)


def chunk(pages):
    """Split each page into overlapping passages of about chunk_words words."""
    size, step = S["chunk_words"], max(1, S["chunk_words"] - S["overlap"])
    out = []
    for page, text in pages:
        words = text.split()
        if not words:
            continue
        for start in range(0, len(words), step):
            piece = words[start:start + size]
            out.append((page, " ".join(piece)))
            if start + size >= len(words):
                break
    return out


def vault_files():
    for p in sorted(S["vault"].rglob("*")):
        rel = p.relative_to(S["vault"])
        if any(part.startswith((".", "_")) for part in rel.parts) or not p.is_file():
            continue
        yield p, str(rel)


# ---------------------------------------------------------------------
# ingest
# ---------------------------------------------------------------------
def cmd_ingest(args):
    if not S["vault"].is_dir():
        sys.exit(f"STOP: no vault at {S['vault']}. Run: bash vault_setup.sh")
    con = open_db()
    model = meta(con, "embed_model")
    if model and model != S["embed_model"]:
        sys.exit(f"STOP: the index was built with {model}, standard.sh now says "
                 f"{S['embed_model']}. Vectors from different models don't mix.\n"
                 f"Run: uv run memory.py rebuild")
    meta(con, "embed_model", S["embed_model"])
    if meta(con, "chunker") not in (None, CHUNKER_VERSION):
        print(f"The way files are read changed (v{meta(con, 'chunker')} -> v{CHUNKER_VERSION}). "
              f"Re-reading every file.\n")
        con.execute("UPDATE files SET sha256=''")
    meta(con, "chunker", CHUNKER_VERSION)

    t0 = time.time()
    counts = dict(new=0, changed=0, same=0, removed=0, skipped=0, scan=0, chunks=0)
    embed_secs = 0.0
    seen = set()
    known = {r[0]: r[1] for r in con.execute("SELECT path, sha256 FROM files")}

    for p, rel in vault_files():
        seen.add(rel)
        data = p.read_bytes()
        sha = hashlib.sha256(data).hexdigest()
        if known.get(rel) == sha:
            counts["same"] += 1
            continue
        kind = p.suffix.lower()
        drop_file(con, rel)
        if kind not in SUPPORTED:
            con.execute("INSERT INTO files VALUES (?,?,?,?,?,?,?,datetime('now'))",
                        (rel, sha, len(data), kind, "skipped", "file type not read yet", 0))
            counts["skipped"] += 1
            continue
        try:
            pages = read_file(p)
        except Exception as e:
            con.execute("INSERT INTO files VALUES (?,?,?,?,?,?,?,datetime('now'))",
                        (rel, sha, len(data), kind, "error", f"{e.__class__.__name__}: {e}", 0))
            print(f"  error   {rel}: {e}")
            continue
        text_len = sum(len(t.strip()) for _, t in pages)
        if kind == ".pdf" and text_len < 20 * max(1, len(pages)):
            con.execute("INSERT INTO files VALUES (?,?,?,?,?,?,?,datetime('now'))",
                        (rel, sha, len(data), kind, "needs OCR",
                         "a scan: no text layer (image reading comes later)", 0))
            counts["scan"] += 1
            print(f"  scan    {rel} (no text layer, not indexed yet)")
            continue
        pieces = chunk(pages)
        te = time.time()
        vecs = embed([t for _, t in pieces]) if pieces else np.zeros((0, 1), np.float32)
        embed_secs += time.time() - te
        if len(pieces):
            if meta(con, "dim") is None:
                meta(con, "dim", vecs.shape[1])
            for i, ((page, text), v) in enumerate(zip(pieces, vecs)):
                cur = con.execute("INSERT INTO chunks(path, ord, page, text, embedding) "
                                  "VALUES (?,?,?,?,?)", (rel, i, page, text, v.tobytes()))
                con.execute("INSERT INTO chunks_fts(rowid, text) VALUES (?,?)",
                            (cur.lastrowid, text))
        con.execute("INSERT INTO files VALUES (?,?,?,?,?,?,?,datetime('now'))",
                    (rel, sha, len(data), kind, "indexed", "", len(pieces)))
        counts["changed" if rel in known else "new"] += 1
        counts["chunks"] += len(pieces)
        print(f"  {'changed' if rel in known else 'new':<7} {rel} ({len(pieces)} passages)")
        con.commit()

    for rel in set(known) - seen:
        drop_file(con, rel)
        counts["removed"] += 1
        print(f"  removed {rel}")
    meta(con, "last_ingest", time.strftime("%Y-%m-%d %H:%M:%S"))
    con.commit()

    secs = time.time() - t0
    print(f"\nnew {counts['new']} · changed {counts['changed']} · unchanged {counts['same']} · "
          f"removed {counts['removed']} · scans {counts['scan']} · other types {counts['skipped']}")
    print(f"{counts['chunks']} passages embedded in {embed_secs:.1f} s · total {secs:.1f} s · "
          f"model {S['embed_model']}")


# ---------------------------------------------------------------------
# search: meaning (vectors) + exact words (FTS5), merged by rank
# ---------------------------------------------------------------------
def fts_query(q):
    words = re.findall(r"\w+", q.lower())
    return " OR ".join(f'"{w}"' for w in words) if words else None


def search(q, k=5, mode="both", pool=50):
    con = open_db()
    ranks = {}  # chunk id -> {"meaning": rank, "words": rank}
    if mode in ("both", "meaning"):
        rows = con.execute("SELECT id, embedding FROM chunks").fetchall()
        if rows:
            ids = np.array([r[0] for r in rows])
            mat = np.frombuffer(b"".join(r[1] for r in rows), dtype=np.float32)
            mat = mat.reshape(len(rows), -1)
            sims = mat @ embed([q], kind="query")[0]
            for rank, j in enumerate(np.argsort(-sims)[:pool]):
                ranks.setdefault(int(ids[j]), {})["meaning"] = (rank, float(sims[j]))
    if mode in ("both", "words"):
        fq = fts_query(q)
        if fq:
            for rank, (cid,) in enumerate(con.execute(
                    "SELECT rowid FROM chunks_fts WHERE chunks_fts MATCH ? "
                    "ORDER BY bm25(chunks_fts) LIMIT ?", (fq, pool))):
                ranks.setdefault(cid, {})["words"] = (rank, None)
    # Reciprocal rank fusion: good ranks in either list rise to the top
    scored = sorted(((sum(1 / (60 + r[0]) for r in d.values()), cid, d)
                     for cid, d in ranks.items()), reverse=True)[:k]
    out = []
    for score, cid, d in scored:
        path, page, text = con.execute(
            "SELECT path, page, text FROM chunks WHERE id=?", (cid,)).fetchone()
        out.append({"path": path, "page": page, "text": text, "score": score,
                    "found_by": "+".join(sorted(d)),
                    "similarity": d.get("meaning", (None, None))[1]})
    return out


def snippet(text, q, width=220):
    words = re.findall(r"\w+", q.lower())
    low = text.lower()
    hits = [low.find(w) for w in words if low.find(w) >= 0]
    start = max(0, min(hits) - 60) if hits else 0
    s = text[start:start + width]
    return ("…" if start else "") + s + ("…" if start + width < len(text) else "")


def cmd_search(args):
    t0 = time.time()
    res = search(args.query, k=args.k, mode=args.mode)
    if not res:
        print("Nothing found in memory.")
        return
    for i, r in enumerate(res, 1):
        where = r["path"] + (f" · p.{r['page']}" if r["path"].lower().endswith(".pdf") else "")
        sim = f" · similarity {r['similarity']:.2f}" if r["similarity"] is not None else ""
        print(f"{i}. {where}  [{r['found_by']}{sim}]")
        print(f"   {snippet(r['text'], args.query)}\n")
    print(f"({time.time() - t0:.2f} s)")


# ---------------------------------------------------------------------
# ask: the main model answers from the passages, and cites them
# ---------------------------------------------------------------------
ASK_RULES = """You answer questions about my own files, using ONLY the numbered passages below.
- After every fact, cite the passage it came from, like [2]. No citation, no fact.
- Copy numbers, dates and names exactly as written. Do not round or convert them.
- If you add or compare numbers, show the working, so I can check it.
- If the passages don't answer the question, say "Not in my memory." and name what kind of file would answer it.
- Never write out a full account, card or ID number. Use the last 4 digits at most.
- Be brief: answer first, then at most a few lines of detail.
- You don't decide for me. If the question asks what I should do, lay out the facts and the choice."""


def cmd_ask(args):
    t0 = time.time()
    hits = search(args.question, k=args.k)
    t_search = time.time() - t0
    if not hits:
        print("Not in my memory. (Nothing indexed matches. Add files to the vault, then ingest.)")
        return
    blocks, sources = [], []
    for i, h in enumerate(hits, 1):
        where = h["path"] + (f" p.{h['page']}" if h["path"].lower().endswith(".pdf") else "")
        sources.append(where)
        blocks.append(f"[{i}] ({where})\n{h['text']}")
    prompt = (f"{ASK_RULES}\n\nPASSAGES\n\n" + "\n\n".join(blocks) +
              f"\n\nQUESTION: {args.question}")
    body = {"model": S["main_model"], "stream": True, "think": args.think,
            "options": {"temperature": 0.2, "num_ctx": S["context"]},
            "messages": [{"role": "user", "content": prompt}]}
    req = urllib.request.Request(f"{S['model_url']}/api/chat", data=json.dumps(body).encode(),
                                 headers={"Content-Type": "application/json"})
    t1, first, n_tokens = time.time(), None, 0
    try:
        with urllib.request.urlopen(req, timeout=600) as r:
            for line in r:
                if not line.strip():
                    continue
                msg = json.loads(line)
                piece = msg.get("message", {}).get("content", "")
                if piece:
                    first = first or time.time()
                    print(piece, end="", flush=True)
                if msg.get("done"):
                    n_tokens = msg.get("eval_count", 0)
    except Exception as e:
        sys.exit(f"\nSTOP: could not reach {S['main_model']} at {S['model_url']} "
                 f"({e.__class__.__name__}: {e})")
    print("\n\nSources")
    for i, src in enumerate(sources, 1):
        print(f"  [{i}] {src}")
    wait = (first - t1) if first else 0
    print(f"\n(search {t_search:.2f} s · first word {wait:.1f} s · answer {time.time() - t1:.1f} s"
          f"{f' · {n_tokens} tokens' if n_tokens else ''} · {S['main_model']})")


# ---------------------------------------------------------------------
# status, rebuild
# ---------------------------------------------------------------------
def cmd_status(args):
    print(f"Vault:  {S['vault']}  {'(found)' if S['vault'].is_dir() else '(MISSING)'}")
    print(f"Index:  {S['index']}  {'(%.1f MB)' % (S['index'].stat().st_size / 1e6) if S['index'].exists() else '(not built)'}")
    print(f"Model:  {S['embed_model']} (search) · {S['main_model']} (answers) at {S['model_url']}")
    try:
        with urllib.request.urlopen(f"{S['model_url']}/api/tags", timeout=3) as r:
            names = [m["name"] for m in json.loads(r.read())["models"]]
        ok = any(n == S["embed_model"] or n.split(":")[0] == S["embed_model"] for n in names)
        print(f"        server reachable · embedding model {'present' if ok else 'NOT PULLED'}")
    except Exception:
        print("        server NOT reachable")
    if not S["index"].exists():
        return
    con = open_db()
    print(f"Built with: {meta(con, 'embed_model')} · {meta(con, 'dim')} dimensions · "
          f"last ingest {meta(con, 'last_ingest')}")
    for status, n, c in con.execute(
            "SELECT status, count(*), sum(chunks) FROM files GROUP BY status ORDER BY status"):
        print(f"  {status:<10} {n:>5} files  {c or 0:>6} passages")
    for path, status, note in con.execute(
            "SELECT path, status, note FROM files WHERE status!='indexed' ORDER BY path LIMIT 20"):
        print(f"    {status:<10} {path}  {note}")


def cmd_rebuild(args):
    for f in S["index"].parent.glob(S["index"].name + "*"):
        f.unlink()
    print(f"Deleted {S['index']}. Building again from the files.\n")
    cmd_ingest(args)


def main():
    p = argparse.ArgumentParser(description="My personal memory")
    sub = p.add_subparsers(dest="cmd", required=True)
    sub.add_parser("ingest", help="index new and changed files").set_defaults(f=cmd_ingest)
    s = sub.add_parser("search", help="find passages")
    s.add_argument("query")
    s.add_argument("-k", type=int, default=5, help="how many results (default 5)")
    s.add_argument("--mode", choices=["both", "meaning", "words"], default="both",
                   help="meaning only, exact words only, or both (default)")
    s.set_defaults(f=cmd_search)
    a_ = sub.add_parser("ask", help="answer a question from my files, with sources")
    a_.add_argument("question")
    a_.add_argument("-k", type=int, default=8, help="how many passages to read (default 8)")
    a_.add_argument("--think", action="store_true", help="let the model think first (slower)")
    a_.set_defaults(f=cmd_ask)
    sub.add_parser("status", help="what's indexed").set_defaults(f=cmd_status)
    sub.add_parser("rebuild", help="delete the index and rebuild").set_defaults(f=cmd_rebuild)
    a = p.parse_args()
    try:
        a.f(a)
    except BrokenPipeError:  # output piped into head, less, etc.
        sys.stderr.close()


if __name__ == "__main__":
    main()
