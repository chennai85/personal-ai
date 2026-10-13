# /// script
# requires-python = ">=3.11"
# dependencies = ["reportlab>=4", "python-docx>=1.1"]
# ///
"""Make made-up sample files of every supported type, for testing. No real data."""
import sys
from email.message import EmailMessage
from pathlib import Path

import docx
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

out = Path(sys.argv[1])
(out / "inbox").mkdir(parents=True, exist_ok=True)

# 1. A two-page PDF with text
c = canvas.Canvas(str(out / "inbox" / "electric-bill-sept.pdf"), pagesize=letter)
c.drawString(72, 720, "Sample Power Co. Statement for September")
c.drawString(72, 700, "Amount due: $142.17. Payment due by October 28.")
c.showPage()
c.drawString(72, 720, "Usage was 812 kWh, up 4 percent from August.")
c.save()

# 1b. A statement whose header repeats on every page; the total is on page 1 only
c = canvas.Canvas(str(out / "inbox" / "year-end-summary.pdf"), pagesize=letter)
for n, body in enumerate(["Total spending for 2025: $12,345.67 across all categories.",
                          "Dining: $2,100.00. Groceries: $3,400.00.",
                          "Travel: $4,800.00. Golf green fees: $2,045.67."], 1):
    c.drawString(72, 750, "Sample Bank 2025 Year-End Summary")
    c.drawString(72, 735, "Includes charges from January 1 through December 31, 2025")
    c.drawString(72, 700, body)
    c.drawString(72, 60, f"Page {n} of 3")
    c.showPage()
c.save()

# 2. A "scan": a PDF with only a picture, no text layer
c = canvas.Canvas(str(out / "inbox" / "scanned-receipt.pdf"), pagesize=letter)
c.rect(72, 600, 300, 100, fill=1)
c.save()

# 3. A Word document with a table
d = docx.Document()
d.add_heading("Car insurance renewal", 1)
d.add_paragraph("The auto policy renews on November 15. Premium for six months is $690.")
t = d.add_table(rows=2, cols=2)
t.cell(0, 0).text, t.cell(0, 1).text = "Coverage", "Limit"
t.cell(1, 0).text, t.cell(1, 1).text = "Collision", "$500 deductible"
d.save(out / "inbox" / "auto-renewal.docx")

# 4. An email
m = EmailMessage()
m["From"], m["To"], m["Subject"] = "clinic@example.com", "me@example.com", "Appointment reminder"
m["Date"] = "Mon, 12 Oct 2026 09:00:00 -0400"
m.set_content("Your physician appointment is Thursday at 10:30 am. Bring your insurance card.")
(out / "inbox" / "reminder.eml").write_bytes(bytes(m))

# 5. A note in Markdown
(out / "inbox" / "trip-ideas.md").write_text(
    "# Vacation ideas\n\nScotland in May: St Andrews and the Highlands.\n"
    "Japan in autumn: Kyoto temples, leaves turning.\n")

# 6. A file type we don't read yet
(out / "inbox" / "photo.heic").write_bytes(b"\x00\x01binary")
print(f"samples written to {out}/inbox")
